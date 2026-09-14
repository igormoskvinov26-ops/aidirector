"""Журнал обзвона: файл, который переживает переустановку.

Проверяется ровно то, ради чего журнал заведён: строка попадает в файл при
каждом звонке, повреждённый файл не стирает будущие записи, а журнал из
прошлой установки вливается без дублей.
"""

from datetime import datetime
from zoneinfo import ZoneInfo

import pytest
import pytest_asyncio
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.config import settings
from app.database import Base
from app.models.models import Client, ContactAttempt, ContactTask
from app.services import call_journal
from app.services.client_base import record_outcome


@pytest.fixture(autouse=True)
def journal_dir(tmp_path, monkeypatch):
    """Журнал пишется во временный каталог, а не в output проекта."""
    monkeypatch.setattr(settings, "output_dir", tmp_path)
    return tmp_path


@pytest_asyncio.fixture
async def session() -> AsyncSession:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    maker = async_sessionmaker(engine, expire_on_commit=False)
    async with maker() as s:
        yield s
    await engine.dispose()


async def _make_task(session: AsyncSession, name: str = "Иван Петров") -> int:
    client = Client(id=1, yclients_id=1001, name=name, phone="+79990000000")
    session.add(client)
    await session.flush()
    task = ContactTask(
        client_id=client.id,
        group_code="risk",
        priority=1,
        due_date=datetime.now().date(),
        status="open",
        script_version="v1",
    )
    session.add(task)
    await session.commit()
    return task.id


def _rows() -> list[tuple]:
    book = load_workbook(call_journal.journal_path())
    return list(book.active.iter_rows(values_only=True))


@pytest.mark.asyncio
async def test_outcome_writes_a_row(session: AsyncSession) -> None:
    task_id = await _make_task(session)

    result = await record_outcome(session, task_id, "booked", "phone", "перезвонил сам", "Катя")

    assert result["ok"] is True
    rows = _rows()
    assert rows[0] == call_journal.COLUMNS
    assert len(rows) == 2
    line = dict(zip(call_journal.COLUMNS, rows[1], strict=True))
    assert line["Клиент"] == "Иван Петров"
    assert line["Телефон"] == "+79990000000"
    assert line["Сегмент"] == "risk"
    # В файл идут слова, а не коды: его открывает администратор, не программа.
    assert line["Результат"] == "записался"
    assert line["Канал"] == "телефон"
    assert line["Комментарий"] == "перезвонил сам"
    assert line["Кто звонил"] == "Катя"


@pytest.mark.asyncio
async def test_rows_accumulate(session: AsyncSession) -> None:
    task_id = await _make_task(session)
    await record_outcome(session, task_id, "no_answer", "phone", None, "Катя")
    await record_outcome(session, task_id, "no_booking", "phone", "дорого", "Катя")

    assert len(_rows()) == 3


@pytest.mark.asyncio
async def test_broken_file_is_moved_aside_not_lost(session: AsyncSession, journal_dir) -> None:
    """Повреждённый файл не должен ни ронять звонок, ни исчезать."""
    call_journal.journal_path().write_bytes(b"not an xlsx at all")
    task_id = await _make_task(session)

    await record_outcome(session, task_id, "booked", "phone", None, "Катя")

    assert len(_rows()) == 2
    broken = list(journal_dir.glob("обзвон-повреждён-*.xlsx"))
    assert len(broken) == 1
    assert broken[0].read_bytes() == b"not an xlsx at all"


@pytest.mark.asyncio
async def test_write_failure_does_not_break_the_call(session: AsyncSession, monkeypatch) -> None:
    """Если файл не пишется, смена всё равно должна продолжаться."""
    monkeypatch.setattr(
        call_journal, "_open_or_create", lambda: (_ for _ in ()).throw(OSError("диск полон"))
    )
    task_id = await _make_task(session)

    result = await record_outcome(session, task_id, "booked", "phone", None, "Катя")

    assert result["ok"] is True
    attempts = (await session.execute(ContactAttempt.__table__.select())).all()
    assert len(attempts) == 1


def test_time_is_moscow_not_container(journal_dir) -> None:
    """Контейнер живёт в UTC, а администратор сверяет строки со своей сменой."""
    utc_evening = datetime(2026, 9, 14, 21, 30, tzinfo=ZoneInfo("UTC"))

    call_journal.append_row(
        when=call_journal._to_moscow(utc_evening),
        client_name="Иван",
        phone="+79990000000",
        segment="risk",
        outcome="booked",
        channel="phone",
        comment=None,
        actor="Катя",
    )

    line = dict(zip(call_journal.COLUMNS, _rows()[1], strict=True))
    assert line["Время"] == "00:30"
    assert line["Дата"] == "15.09.2026"


def _write_journal(path, rows: list[list]) -> None:
    book = Workbook()
    book.active.append(list(call_journal.COLUMNS))
    for row in rows:
        book.active.append(row)
    book.save(path)


def test_import_merges_without_duplicates(journal_dir) -> None:
    old = ["14.09.2026", "12:00", "Иван", "+79990000000",
           "risk", "записался", "телефон", "", "Катя"]
    new = ["14.09.2026", "13:00", "Пётр", "+79991111111",
           "late", "не ответил", "телефон", "", "Катя"]
    _write_journal(call_journal.journal_path(), [old])
    source = journal_dir / "из-салона.xlsx"
    _write_journal(source, [old, new])

    report = call_journal.merge_journal(source)

    assert report == {"прочитано": 2, "добавлено": 1, "пропущено_дублей": 1}
    assert len(_rows()) == 3


def test_import_rejects_a_foreign_file(journal_dir) -> None:
    """Чужая таблица не должна молча попасть в журнал."""
    source = journal_dir / "чужая.xlsx"
    book = Workbook()
    book.active.append(["Товар", "Цена"])
    book.active.append(["Шампунь", 900])
    book.save(source)

    with pytest.raises(ValueError, match="не журнал обзвона"):
        call_journal.merge_journal(source)


@pytest.mark.asyncio
async def test_rebuild_keeps_the_previous_file(session: AsyncSession, journal_dir) -> None:
    """Пересборка из базы — операция по просьбе, и прежний файл она не съедает."""
    task_id = await _make_task(session)
    await record_outcome(session, task_id, "booked", "phone", None, "Катя")

    written = await call_journal.rebuild_from_database(session)

    assert written == 1
    assert len(list(journal_dir.glob("обзвон-до-пересборки-*.xlsx"))) == 1
    assert len(_rows()) == 2


# ── Кто что может ─────────────────────────────────────────────────────────── #


def _app_client():
    """Приложение с настоящими маршрутами журнала и настоящей авторизацией."""
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.api.routes.client_base import router
    from app.main import BasicAuthMiddleware

    app = FastAPI()
    app.add_middleware(BasicAuthMiddleware)
    app.include_router(router)
    return TestClient(app)


def _auth(login: str, password: str) -> dict[str, str]:
    import base64

    return {"Authorization": "Basic " + base64.b64encode(f"{login}:{password}".encode()).decode()}


def test_operator_downloads_but_does_not_import(journal_dir) -> None:
    """Администратор работает журналом, но не переписывает историю салона."""
    call_journal.append_row(
        when=datetime(2026, 9, 14, 12, 0),
        client_name="Иван",
        phone="+79990000000",
        segment="risk",
        outcome="booked",
        channel="phone",
        comment=None,
        actor="Катя",
    )
    client = _app_client()
    operator = _auth(settings.operator_login, settings.operator_password)

    assert client.get("/api/client-base/journal", headers=operator).status_code == 200

    upload = {"file": ("обзвон.xlsx", b"whatever", "application/vnd.ms-excel")}
    assert client.post(
        "/api/client-base/journal/import", headers=operator, files=upload
    ).status_code == 403


def test_owner_imports(journal_dir) -> None:
    source = journal_dir / "из-салона.xlsx"
    _write_journal(source, [["14.09.2026", "12:00", "Иван", "+79990000000",
                             "risk", "записался", "телефон", "", "Катя"]])
    client = _app_client()
    owner = _auth(settings.owner_login, settings.owner_password)

    with source.open("rb") as handle:
        response = client.post(
            "/api/client-base/journal/import",
            headers=owner,
            files={"file": ("обзвон.xlsx", handle, "application/vnd.ms-excel")},
        )

    assert response.status_code == 200
    assert response.json()["добавлено"] == 1


def test_empty_journal_says_so(journal_dir) -> None:
    client = _app_client()
    owner = _auth(settings.owner_login, settings.owner_password)

    assert client.get("/api/client-base/journal", headers=owner).status_code == 404
    assert client.get("/api/client-base/journal/status", headers=owner).json()["exists"] is False
