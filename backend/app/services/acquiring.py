"""Отчёт банка по эквайрингу: разбор файла, хранение, суммы по дням.

Формат банка заранее неизвестен, поэтому колонки ищутся по заголовкам
(«дата операции», «сумма», «комиссия», «к зачислению»). Если дата или сумма
не нашлись, загрузка отклоняется с перечнем найденных заголовков: молча
читать не ту колонку хуже, чем попросить образец файла.
"""

import csv
import io
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.models import AcquiringRow

ЗАГОЛОВКОВ_ИСКАТЬ = 40  # строк сверху, где может быть шапка таблицы
CENT = Decimal("0.01")


def _текст(v: object) -> str:
    return " ".join(str(v).split()).lower() if v is not None else ""


def _таблица(content: bytes, filename: str) -> list[list[object]]:
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows: list[list[object]] = []
        for ws in wb.worksheets:
            rows.extend([list(r) for r in ws.iter_rows(values_only=True)])
        return rows
    if name.endswith(".csv") or name.endswith(".txt"):
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("Не удалось прочитать кодировку файла.")
        sep = ";" if text.count(";") >= text.count(",") else ","
        return [list(r) for r in csv.reader(io.StringIO(text), delimiter=sep)]
    if name.endswith(".xls"):
        raise ValueError("Формат .xls не поддерживается: сохраните отчёт как .xlsx или .csv.")
    raise ValueError("Нужен файл Excel (.xlsx) или CSV.")


def _дата(v: object) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()[:19]
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
                "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _сумма(v: object) -> Decimal | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int | float):
        return Decimal(str(v)).quantize(CENT)
    s = str(v).replace("\xa0", "").replace(" ", "").replace(",", ".").replace("−", "-")
    if not s:
        return None
    try:
        return Decimal(s).quantize(CENT)
    except InvalidOperation:
        return None


def _колонка(шапка: list[str], *, нужно: tuple[str, ...], предпочесть: tuple[str, ...] = (),
             нельзя: tuple[str, ...] = ()) -> int | None:
    подходят = [i for i, h in enumerate(шапка)
                if h and any(w in h for w in нужно) and not any(w in h for w in нельзя)]
    for i in подходят:
        if any(w in шапка[i] for w in предпочесть):
            return i
    return подходят[0] if подходят else None


def разобрать(content: bytes, filename: str) -> dict:
    """Строки отчёта: [{date, amount, fee, net}] и сведения о найденных колонках."""
    таблица = _таблица(content, filename)
    найдено: list[str] = []
    for номер, строка in enumerate(таблица[:ЗАГОЛОВКОВ_ИСКАТЬ]):
        шапка = [_текст(c) for c in строка]
        if sum(1 for h in шапка if "дата" in h) and any(
            w in h for h in шапка for w in ("сумма", "оборот")
        ):
            найдено = [str(c) for c in строка if c not in (None, "")]
            break
    else:
        raise ValueError("Не нашёл строку с заголовками (дата и сумма). Пришлите образец файла.")

    d = _колонка(шапка, нужно=("дата",), предпочесть=("операц", "транзакц", "оплат", "покупк"),
                 нельзя=("зачисл", "перечисл", "расчёт", "расчет"))
    a = _колонка(шапка, нужно=("сумма", "оборот"),
                 предпочесть=("операц", "покупк", "оплат", "оборот", "транзакц"),
                 нельзя=("комисс", "зачисл", "перечисл", "выплат"))
    f = _колонка(шапка, нужно=("комисс",), предпочесть=("сумма", "руб"),
                 нельзя=("%", "процент", "ставк"))
    n = _колонка(шапка, нужно=("зачисл", "перечисл", "к выплате"))
    t = _колонка(шапка, нужно=("тип операции", "вид операции"))
    итог_шапки = None
    for строка in таблица[:номер]:
        if len(строка) > 1 and _текст(строка[0]).startswith("сумма транзакций") \
                and "без" not in _текст(строка[0]):
            итог_шапки = _сумма(строка[1])
    if d is None or a is None:
        raise ValueError(
            "Не распознал колонки даты и суммы. Найдены заголовки: " + ", ".join(найдено)
            + ". Пришлите образец файла."
        )

    rows: list[dict] = []
    пропущено = 0
    for строка in таблица[номер + 1:]:
        if len(строка) <= max(d, a):
            пропущено += 1
            continue
        день, сумма = _дата(строка[d]), _сумма(строка[a])
        if день is None or сумма is None:
            пропущено += 1  # «Итого», пустые и служебные строки
            continue
        комиссия = _сумма(строка[f]) if f is not None and f < len(строка) else None
        комиссия = abs(комиссия) if комиссия is not None else Decimal(0)
        зачислено = _сумма(строка[n]) if n is not None and n < len(строка) else None
        вид = _текст(строка[t]) if t is not None and t < len(строка) else ""
        if "дебет" in вид or "возврат" in вид:  # возврат клиенту уменьшает оборот
            сумма, комиссия = -abs(сумма), -комиссия
            зачислено = None
        rows.append({"date": день, "amount": сумма, "fee": комиссия,
                     "net": зачислено if зачислено is not None else (сумма - комиссия)})
    if not rows:
        raise ValueError("В файле нет строк с датой и суммой. Пришлите образец файла.")
    сумма_строк = sum((r["amount"] for r in rows), Decimal(0))
    return {
        "rows": rows, "skipped": пропущено, "header_total": итог_шапки,
        "matches_header": None if итог_шапки is None else abs(итог_шапки - сумма_строк) <= CENT,
        "columns": {"date": шапка[d], "amount": шапка[a],
                    "fee": шапка[f] if f is not None else None,
                    "net": шапка[n] if n is not None else None},
    }


async def сохранить(
    session: AsyncSession, content: bytes, filename: str, день: date | None = None
) -> dict:
    """Разобрать и записать. Даты из файла заменяют ранее загруженные за эти же даты.

    день — отчёт загружают «за этот день» с графика сверки: если такого дня в
    файле нет, это почти наверняка не тот файл, и записывать его нельзя."""
    разбор = разобрать(content, filename)
    rows = разбор["rows"]
    даты = {r["date"] for r in rows}
    if день is not None and день not in даты:
        raise ValueError(
            f"В файле нет операций за {день:%d.%m.%Y}: в нём {min(даты):%d.%m}–{max(даты):%d.%m}. "
            "Проверьте, тот ли это отчёт."
        )
    await session.execute(delete(AcquiringRow).where(AcquiringRow.op_date.in_(даты)))
    for r in rows:
        session.add(AcquiringRow(op_date=r["date"], amount=r["amount"], fee=r["fee"],
                                 net=r["net"], source_file=filename[:255]))
    await session.commit()
    return {
        "rows": len(rows), "skipped": разбор["skipped"], "days": len(даты),
        "date_from": min(даты).isoformat(), "date_to": max(даты).isoformat(),
        "amount": float(sum(r["amount"] for r in rows)),
        "fee": float(sum(r["fee"] for r in rows)),
        "net": float(sum(r["net"] for r in rows)),
        "columns": разбор["columns"],
        "header_total": float(разбор["header_total"]) if разбор["header_total"] is not None else None,
        "matches_header": разбор["matches_header"],
    }


async def по_дням(session: AsyncSession, start: date, end: date) -> dict[date, dict[str, Decimal]]:
    """Суммы отчёта банка по дням включительно: оборот, комиссия, зачисление."""
    res = await session.execute(
        select(AcquiringRow).where(AcquiringRow.op_date >= start, AcquiringRow.op_date <= end)
    )
    days: dict[date, dict[str, Decimal]] = defaultdict(
        lambda: {"amount": Decimal(0), "fee": Decimal(0), "net": Decimal(0)}
    )
    for r in res.scalars():
        days[r.op_date]["amount"] += r.amount
        days[r.op_date]["fee"] += r.fee
        days[r.op_date]["net"] += r.net
    return dict(days)
